from pathlib import Path
from fastapi import HTTPException
import html as html_mod


def pdf_to_word(input_path: Path, output_path: Path) -> Path:
    """Converte PDF in DOCX usando pdf2docx."""
    try:
        from pdf2docx import Converter
        cv = Converter(str(input_path))
        cv.convert(str(output_path), start=0, end=None)
        cv.close()
        return output_path
    except ImportError:
        raise HTTPException(status_code=500, detail="pdf2docx non installato.")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Errore PDF→Word: {str(e)}")


def _docx_to_pdf_reportlab(input_path: Path, output_path: Path) -> Path:
    """
    Converte DOCX in PDF usando python-docx + reportlab.
    Puro Python, zero dipendenze di sistema.
    """
    from docx import Document
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.enums import TA_LEFT, TA_CENTER, TA_RIGHT

    try:
        doc = Document(str(input_path))
        pdf_doc = SimpleDocTemplate(
            str(output_path),
            pagesize=A4,
            leftMargin=2*cm, rightMargin=2*cm,
            topMargin=2*cm, bottomMargin=2*cm
        )

        styles = getSampleStyleSheet()
        story = []

        style_normal = ParagraphStyle('Normal', parent=styles['Normal'], fontSize=11, leading=16)
        style_h1 = ParagraphStyle('H1', parent=styles['Heading1'], fontSize=18, spaceAfter=8)
        style_h2 = ParagraphStyle('H2', parent=styles['Heading2'], fontSize=15, spaceAfter=6)
        style_h3 = ParagraphStyle('H3', parent=styles['Heading3'], fontSize=13, spaceAfter=4)

        for para in doc.paragraphs:
            text = para.text.strip()
            if not text:
                story.append(Spacer(1, 6))
                continue

            safe = html_mod.escape(text)
            style_name = para.style.name.lower()

            if 'heading 1' in style_name:
                story.append(Paragraph(safe, style_h1))
            elif 'heading 2' in style_name:
                story.append(Paragraph(safe, style_h2))
            elif 'heading 3' in style_name:
                story.append(Paragraph(safe, style_h3))
            else:
                # Applica bold/italic inline
                parts = []
                for run in para.runs:
                    r = html_mod.escape(run.text)
                    if not r:
                        continue
                    if run.bold and run.italic:
                        r = f"<b><i>{r}</i></b>"
                    elif run.bold:
                        r = f"<b>{r}</b>"
                    elif run.italic:
                        r = f"<i>{r}</i>"
                    parts.append(r)
                content = ''.join(parts) if parts else safe
                story.append(Paragraph(content, style_normal))
                story.append(Spacer(1, 2))

        # Tabelle
        for table in doc.tables:
            data = []
            for row in table.rows:
                data.append([html_mod.escape(cell.text) for cell in row.cells])
            if data:
                t = Table(data, hAlign='LEFT')
                t.setStyle(TableStyle([
                    ('GRID', (0,0), (-1,-1), 0.5, colors.grey),
                    ('BACKGROUND', (0,0), (-1,0), colors.lightgrey),
                    ('FONTSIZE', (0,0), (-1,-1), 10),
                    ('PADDING', (0,0), (-1,-1), 4),
                ]))
                story.append(t)
                story.append(Spacer(1, 8))

        if not story:
            story.append(Paragraph("Documento vuoto", style_normal))

        pdf_doc.build(story)
        return output_path

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Errore conversione Word→PDF: {str(e)}")


def _xlsx_to_pdf_reportlab(input_path: Path, output_path: Path) -> Path:
    """Converte XLSX in PDF usando openpyxl + reportlab."""
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    try:
        import openpyxl
        wb = openpyxl.load_workbook(str(input_path), data_only=True)
        ws = wb.active

        data = []
        for row in ws.iter_rows(values_only=True):
            data.append([str(c) if c is not None else '' for c in row])

        if not data:
            raise HTTPException(status_code=400, detail="Foglio Excel vuoto.")

        pdf_doc = SimpleDocTemplate(str(output_path), pagesize=landscape(A4),
                                     leftMargin=1*cm, rightMargin=1*cm,
                                     topMargin=1.5*cm, bottomMargin=1.5*cm)
        styles = getSampleStyleSheet()
        story = []
        story.append(Paragraph(ws.title or "Foglio", styles['Heading2']))
        story.append(Spacer(1, 8))

        t = Table(data, repeatRows=1)
        t.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 0.4, colors.grey),
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e0e0e0')),
            ('FONTSIZE', (0,0), (-1,-1), 8),
            ('PADDING', (0,0), (-1,-1), 3),
            ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ]))
        story.append(t)
        pdf_doc.build(story)
        return output_path

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Errore Excel→PDF: {str(e)}")


def _pptx_to_pdf_reportlab(input_path: Path, output_path: Path) -> Path:
    """Converte PPTX in PDF usando python-pptx + reportlab."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    try:
        from pptx import Presentation
        prs = Presentation(str(input_path))
        pdf_doc = SimpleDocTemplate(str(output_path), pagesize=A4,
                                     leftMargin=2*cm, rightMargin=2*cm,
                                     topMargin=2*cm, bottomMargin=2*cm)
        styles = getSampleStyleSheet()
        style_slide = ParagraphStyle('Slide', fontSize=9, textColor=colors.grey)
        story = []

        for i, slide in enumerate(prs.slides, 1):
            story.append(Paragraph(f"— Slide {i} —", style_slide))
            story.append(Spacer(1, 4))
            for shape in slide.shapes:
                if not hasattr(shape, "text") or not shape.text.strip():
                    continue
                text = html_mod.escape(shape.text.strip())
                if shape == slide.shapes[0]:
                    story.append(Paragraph(text, styles['Heading2']))
                else:
                    story.append(Paragraph(text, styles['Normal']))
                story.append(Spacer(1, 4))
            story.append(HRFlowable(width="100%", thickness=0.5, color=colors.lightgrey))
            story.append(Spacer(1, 12))

        if not story:
            story.append(Paragraph("Presentazione vuota", styles['Normal']))

        pdf_doc.build(story)
        return output_path

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Errore PPTX→PDF: {str(e)}")


def word_to_pdf(input_path: Path, output_path: Path) -> Path:
    return _docx_to_pdf_reportlab(input_path, output_path)


def excel_to_pdf(input_path: Path, output_path: Path) -> Path:
    return _xlsx_to_pdf_reportlab(input_path, output_path)


def pptx_to_pdf(input_path: Path, output_path: Path) -> Path:
    return _pptx_to_pdf_reportlab(input_path, output_path)


def pdf_to_images(input_path: Path, output_dir: Path, fmt: str = "jpg", dpi: int = 150) -> list:
    """Converte ogni pagina PDF in immagine usando PyMuPDF."""
    try:
        import fitz
        doc = fitz.open(str(input_path))
        paths = []
        for i, page in enumerate(doc):
            mat = fitz.Matrix(dpi / 72, dpi / 72)
            pix = page.get_pixmap(matrix=mat)
            out = output_dir / f"page_{i+1}.{fmt}"
            pix.save(str(out), "jpeg" if fmt.lower() in ("jpg","jpeg") else fmt)
            paths.append(out)
        doc.close()
        return paths
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Errore PDF→Immagini: {str(e)}")
